"""Edit the supplied 41-slide deck in place in memory; write a separate final copy."""
from pathlib import Path
import json, hashlib, re, math
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

ROOT=Path(__file__).resolve().parent; PROJECT=ROOT.parent
R=json.loads((ROOT/'final_results.json').read_text()); HP=json.loads((ROOT/'hyperparameters.json').read_text())
S=json.loads((ROOT/'split_summary.json').read_text()); CAL=json.loads((ROOT/'calibration_metrics.json').read_text())
RAW=R['raw_metrics']; SM=R['smoothed_metrics']; RUN=R['run_id']
REPORT=pd.read_csv(ROOT/'classification_report.csv',index_col=0)
BIO=pd.read_csv(ROOT/'subject_biomarkers.csv'); V=pd.read_csv(ROOT/'biomarker_validation.csv')
CM=pd.read_csv(ROOT/'confusion_matrix_normalized.csv',index_col=0)
LAY=pd.read_csv(ROOT/'model_layers.csv'); CW=pd.read_csv(ROOT/'class_weights.csv')
NAVY='0B2E4F';TEAL='13A6A1';BLUE='2F6FED';GRAY='60788A';LIGHT='EEF5F8';INK='19364D'
prs=Presentation(str(PROJECT/'mid.original_backup.pptx'))
assert len(prs.slides)==41
W=prs.slide_width/Inches(1);H=prs.slide_height/Inches(1)
manifest=[]
def rgb(h):return RGBColor.from_string(h)
def box(sl,x,y,w,h,fill=LIGHT,line=None,rounded=False):
    sp=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x),Inches(y),Inches(w),Inches(h))
    sp.fill.solid();sp.fill.fore_color.rgb=rgb(fill)
    if line:sp.line.color.rgb=rgb(line)
    else:sp.line.fill.background()
    return sp
def text(sl,txt,x,y,w,h,size=19,color=INK,bold=False,align=PP_ALIGN.LEFT,fill=None):
    sp=box(sl,x,y,w,h,fill) if fill else sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=sp.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=Inches(.10 if fill else .02);tf.margin_right=Inches(.10 if fill else .02)
    tf.margin_top=Inches(.05);tf.margin_bottom=Inches(.04);tf.vertical_anchor=MSO_ANCHOR.TOP
    for i,line in enumerate(str(txt).split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.alignment=align
        p.space_after=Pt(9 if '\n' in str(txt) else 0)
        p.font.name='Arial';p.font.size=Pt(size);p.font.color.rgb=rgb(color);p.font.bold=bold
    return sp
def picture(sl,name,x,y,w,h):
    p=ROOT/name
    with Image.open(p) as im:iw,ih=im.size
    ratio=min(w/iw,h/ih);dw=iw*ratio;dh=ih*ratio
    return sl.shapes.add_picture(str(p),Inches(x+(w-dw)/2),Inches(y+(h-dh)/2),width=Inches(dw),height=Inches(dh))
def table(sl,headers,rows,x=.65,y=1.55,w=12.05,h=None,widths=None,size=18,risk_col=None):
    h=h or .58*(len(rows)+1)
    tbl=sl.shapes.add_table(len(rows)+1,len(headers),Inches(x),Inches(y),Inches(w),Inches(h)).table
    if widths:
        for c,width in zip(tbl.columns,widths):c.width=Inches(width)
    for i,row in enumerate([headers]+list(rows)):
        for j,value in enumerate(row):
            cell=tbl.cell(i,j);cell.text=str(value);cell.margin_left=Inches(.09);cell.margin_right=Inches(.08)
            cell.margin_top=Inches(.09);cell.margin_bottom=Inches(.06);cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            cell.fill.solid();cell.fill.fore_color.rgb=rgb(NAVY if i==0 else ('FFFFFF' if i%2 else 'F1F6F8'))
            if i and j==risk_col:
                cell.fill.fore_color.rgb=rgb({'Healthy':'DBF1E9','Mild Risk':'FFF0CF','High Risk':'FADADA'}.get(str(value),'EEF5F8'))
            for p in cell.text_frame.paragraphs:
                p.font.name='Arial';p.font.size=Pt(size);p.font.bold=(i==0);p.font.color.rgb=rgb('FFFFFF' if i==0 else INK)
    return tbl
def bullets(sl,items,x=.8,y=1.45,w=11.8,h=4.9,size=22):
    return text(sl,'\n'.join('• '+str(t) for t in items),x,y,w,h,size)
def banner(sl,txt,y=6.42,color=LIGHT):text(sl,txt,.7,y,11.95,.47,16,fill=color)
def card(sl,title,body,x,y,w,h=1.7):
    box(sl,x,y,w,h,LIGHT,rounded=True);box(sl,x,y,.05,h,TEAL)
    text(sl,title,x+.17,y+.12,w-.34,.4,21,NAVY,True)
    text(sl,body,x+.17,y+.50,w-.34,h-.54,18)
def arrow(sl,x,y,w=.4,h=.3,direction='right'):
    sp=sl.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW if direction=='right' else MSO_SHAPE.DOWN_ARROW,
        Inches(x),Inches(y),Inches(w),Inches(h));sp.fill.solid();sp.fill.fore_color.rgb=rgb(BLUE);sp.line.fill.background()
def flow(sl,labels,y=2.4,height=1.35):
    n=len(labels);gap=.4;bw=(12.05-gap*(n-1))/n
    for i,label in enumerate(labels):
        x=.65+i*(bw+gap);box(sl,x,y,bw,height,LIGHT,rounded=True)
        text(sl,label,x+.06,y+.20,bw-.12,height-.3,18,NAVY,True,PP_ALIGN.CENTER)
        if i<n-1:arrow(sl,x+bw+.06,y+height/2-.12,.29,.24)
def prepare(i,title=None,sources=()):
    sl=prs.slides[i-1]
    # Keep the original full-slide background and the existing title element.
    title_sp=next((sp for sp in sl.shapes if sp.has_text_frame and sp.text.strip()),None)
    keep=[]
    for sp in list(sl.shapes):
        full=(sp.left==0 and sp.top==0 and sp.width>=prs.slide_width-1000 and sp.height==prs.slide_height)
        if full or (title_sp is not None and sp._element is title_sp._element):keep.append(sp)
        else:sp._element.getparent().remove(sp._element)
    if title_sp is None:title_sp=text(sl,title or '',.65,.4,12,.6)
    title=title or title_sp.text
    title_sp.left=Inches(.65);title_sp.top=Inches(.38);title_sp.width=Inches(12);title_sp.height=Inches(.7)
    tf=title_sp.text_frame;tf.clear();tf.word_wrap=False;tf.margin_left=0;tf.margin_top=0;tf.margin_bottom=0
    p=tf.paragraphs[0];p.text=title;p.font.name='Arial';p.font.size=Pt(28 if len(title)<66 else 25)
    p.font.color.rgb=rgb(NAVY);p.font.bold=True
    box(sl,.65,1.13,12.05,.035,TEAL)
    footer='Source: '+RUN+' | '+', '.join(sources) if sources else 'Research system — not a medical diagnosis'
    text(sl,footer,.68,7.05,11.8,.25,10.5,GRAY)
    text(sl,f'{i:02d}',12.45,7.02,.4,.28,11,GRAY,align=PP_ALIGN.RIGHT)
    sl.notes_slide.notes_text_frame.text=f"{title}\nRun: {RUN}\nSources: {', '.join(sources)}\nTest recordings: {', '.join(R['test_subjects'])}.\nCheckpoint: {R['checkpoint']}\nCheckpoint SHA256: {R['checkpoint_sha256']}\nAll classification results use identical corrected test epochs.\n"+'\n'.join(R['limitations'])
    manifest.append({'slide':i,'title':title,'sources':list(sources),'run_id':RUN})
    return sl
def pct(v):return f'{v*100:.2f}%'

# Cover: retain supplied design and author/guide/institution text.
sl=prs.slides[0]
for sp in sl.shapes:
    if sp.has_text_frame:
        replacement={'PROJECT PROGRESS REPORT':'FINAL PROJECT REPORT',
                     'Format: Project Progress Report':'Final evaluation: 7 October 2026'}.get(sp.text)
        if replacement:
            runs=sp.text_frame.paragraphs[0].runs
            runs[0].text=replacement
            for run in runs[1:]:run.text=''
        for p in sp.text_frame.paragraphs:
            for run in p.runs:
                if not run.font.name:run.font.name='Arial'
for sp in sl.shapes:
    if sp.has_text_frame and sp.name=='Text 3':
        sp.height=Inches(1.25);sp.width=Inches(11.1)
    if sp.has_text_frame and sp.name=='Text 12':
        sp.width+=Inches(.12);sp.height+=Inches(.05)
sl.notes_slide.notes_text_frame.text='Authors, enrollment numbers, guide and institution are supplied by the project owner. Full project title retained across title and subtitle. '+RUN
manifest.append({'slide':1,'title':'Cover','sources':['user-supplied project metadata']})

sl=prepare(2,sources=['project scope','preprocessing_config.json'])
bullets(sl,['Classify four-channel EEG into Wake, N1, N2, N3 and REM at 30-second resolution.',
    'Automate epoch scoring, expose evidential uncertainty, and decode recording-level stage sequences.',
    'Convert scored sequences into interpretable research sleep profiles.'],y=1.5,h=2.35,size=23)
flow(sl,['EEG','Five-stage\nclassification','Sequence\ndecoding','Research\nprofile'],y=4.0)
banner(sl,'Scope: academic engineering research; no diagnostic or clinical-validation claim.')

sl=prepare(3,sources=['checkpoint_config.json','transition_config.json','biomarker_equations.md'])
for j,(title,body) in enumerate([('01  Sleep-stage classification','Predict five stages using a compact CNN–self-attention–BiLSTM evidential model.'),
    ('02  Sequence refinement','Decode stage probabilities with training-derived transitions, separately for each recording.'),
    ('03  Subject-level profiling','Compute source-defined biomarkers, compare with ground truth and apply research risk rules.')]):
    card(sl,title,body,.7,1.45+j*1.65,11.95,1.45)

sl=prepare(4,sources=['split_summary.json','split_subjects.csv','input_repairs.json'])
flow(sl,['24 recordings\n24 unique SN IDs','16 training\n15,216 epochs','4 validation\n3,619 epochs',f"4 test\n{RAW['test_epochs']:,} epochs"],y=1.6,height=1.2)
totals=pd.read_csv(ROOT/'stage_distribution.csv').groupby('stage')['count'].sum()
text(sl,f"Total evaluated usable epochs: {S['total_usable_epochs']:,}\nOverall counts: "+' · '.join(f'{n} {int(totals[n]):,}' for n in ['Wake','N1','N2','N3','REM']),.8,3.05,11.6,1.05,21)
text(sl,'Selected EEG channels: F4–M1 · C4–M1 · O2–M1 · C3–M2\nTest IDs: '+', '.join(R['test_subjects']),.8,4.3,11.6,.95,21)
text(sl,'SN001 correction: 190 stored epochs → 854 EDF-derived epochs. Original data preserved.\nNo subject/recording ID appears in more than one split.',.8,5.5,11.6,.85,18)
sl.notes_slide.notes_text_frame.text+='\nFull split: '+json.dumps(S['split_subject_ids'])

sl=prepare(5,sources=['preprocessing_config.json','eeg_example_metadata.json'])
picture(sl,'raw_eeg_example.png',.55,1.4,6.1,3.5);picture(sl,'preprocessed_eeg_example.png',6.65,1.4,6.1,3.5)
text(sl,'Channel selection → 50 Hz notch → 0.3–35 Hz bandpass → resample 256→100 Hz',.8,5.05,11.8,.52,21,NAVY,True)
text(sl,'30-second epochs → channel-wise epoch normalization → input shape (4, 3000)\nNo amplitude-based rejection; reject_by_annotation=False in the source.',.8,5.72,11.8,.9,19)

sl=prepare(6,sources=['stage_distribution.csv','class_weights.csv'])
picture(sl,'stage_distribution.png',.7,1.32,11.95,5.15)
banner(sl,'Class weighting: Ntraining / (5 × training class count); computed from training labels only.')

sl=prepare(7,sources=['checkpoint_config.json','model_layers.csv'])
flow(sl,['EEG epoch\n4 × 3000','Conv1D 64\nk7, stride4\n750 × 64','Conv1D 128\nk5, stride4\n188 × 128','Self-attention\n4 heads, key16\n188 × 128'],y=1.55,height=1.45)
flow(sl,['Softplus 5\nevidence\nα = e + 1','Average pool\nDense 128\nDropout .2/.1','BiLSTM 64×2\n188 × 128','Residual + LN\nFFN 128→128\nGELU'],y=3.65,height=1.45)
arrow(sl,11.4,3.12,.35,.40,'down')
for sp in sl.shapes:
    if sp.shape_type==1 and getattr(sp,'auto_shape_type',None)==MSO_SHAPE.RIGHT_ARROW and sp.top>Inches(3.6):sp.rotation=180
banner(sl,f"Hybrid CNN–Transformer/self-attention–BiLSTM | {HP['trainable_parameters']:,} trainable parameters")

sl=prepare(8,sources=['checkpoint_config.json','model_layers.csv'])
for j,(title,body) in enumerate([('CNN downsampling','Convolutional filters learn local patterns; strides reduce 3000 samples to 188 feature positions.'),
    ('Self-attention + feed-forward block','Four attention heads combine feature positions within each individual 30-second epoch.'),
    ('Bidirectional LSTM','64 units in each direction model the downsampled within-epoch sequence.'),
    ('Evidential output','Five nonnegative evidence values define Dirichlet means and total-evidence uncertainty.')]):
    card(sl,title,body,.7,1.45+j*1.27,11.95,1.13)

sl=prepare(9,sources=['calibration_metrics.json','test_uncertainty.csv'])
picture(sl,'uncertainty_correct_incorrect.png',.55,1.3,8.45,5.25)
text(sl,f"Actual output\nαk = evidencek + 1\npk = αk / Σα\nu = 5 / Σα\n\nError-detection AUROC\n{CAL['error_detection_AUROC']:.3f}",9.15,1.75,3.5,3.7,21)
banner(sl,'Higher uncertainty is useful only if it separates errors; this is evaluated, not assumed.')

sl=prepare(10,sources=['training_history.csv','hyperparameters.json'])
picture(sl,'training_accuracy.png',.55,1.35,6.1,3.35);picture(sl,'training_loss.png',6.65,1.35,6.1,3.35)
table(sl,['Setting','Verified value','Setting','Verified value'],[
    ['Optimizer','Adam','Learning rate','1 × 10⁻⁴'],['Batch size','128','Epochs','10; selected 9'],
    ['Early stopping','Patience 5; not triggered','Training hardware','Not preserved in run log']],y=4.85,h=1.65,size=17,widths=[2.05,3.0,2.5,4.5])

sl=prepare(11,'Sleep-Stage Classification Performance',sources=['overall_metrics.json','test_epoch_predictions.csv'])
metrics=[('Accuracy',pct(RAW['accuracy'])),('Balanced accuracy',pct(RAW['balanced_accuracy'])),('Macro F1',pct(RAW['macro_f1'])),('Weighted F1',pct(RAW['weighted_f1'])),("Cohen’s kappa",f"{RAW['cohen_kappa']:.4f}")]
for j,(label,value) in enumerate(metrics):
    x=.75+(j%3)*4.13;y=1.6+(j//3)*2.05
    box(sl,x,y,3.9,1.7,LIGHT,rounded=True);text(sl,label,x+.18,y+.17,3.5,.45,20,NAVY)
    text(sl,value,x+.18,y+.68,3.5,.8,36,TEAL,True)
text(sl,f"Raw evidential model\n{RAW['test_epochs']:,} epochs · four held-out test IDs",9.12,4.0,3.25,1.2,21,NAVY,True)
banner(sl,'Final checkpoint, corrected test inputs, one run. Subject-disjoint fit; no external validation.')

sl=prepare(12,sources=['confusion_matrix_normalized.csv','classification_report.csv'])
picture(sl,'confusion_matrix_normalized.png',.5,1.25,7.7,5.62)
best=REPORT.loc[['Wake','N1','N2','N3','REM'],'recall'].idxmax();worst=REPORT.loc[['Wake','N1','N2','N3','REM'],'recall'].idxmin()
off=CM.to_numpy().copy();np=__import__('numpy');np.fill_diagonal(off,0);ii,jj=np.unravel_index(off.argmax(),off.shape)
bullets(sl,[f"Highest recall: {best} ({pct(REPORT.loc[best,'recall'])}).",f"Lowest recall: {worst} ({pct(REPORT.loc[worst,'recall'])}).",
    f"Largest off-diagonal share: {CM.index[ii]} → {CM.columns[jj]} ({off[ii,jj]*100:.2f}%)."],x=8.4,y=1.8,w=4.1,h=3.7,size=21)

sl=prepare(13,sources=['classification_report.csv'])
picture(sl,'class_f1.png',.65,1.35,8.0,4.9)
table(sl,['Stage','Precision','Recall'],[[n,f"{REPORT.loc[n,'precision']:.3f}",f"{REPORT.loc[n,'recall']:.3f}"] for n in ['Wake','N1','N2','N3','REM']],x=8.75,y=1.85,w=3.9,h=3.6,size=16,widths=[1.05,1.43,1.42])
banner(sl,f"Lowest class F1: {REPORT.loc[['Wake','N1','N2','N3','REM'],'f1-score'].idxmin()}; all values come from the raw-model condition.")

sl=prepare(14,sources=['calibration_metrics.json','calibration_bins.csv'])
picture(sl,'uncertainty_correct_incorrect.png',.55,1.35,6.15,3.75);picture(sl,'reliability_diagram.png',6.8,1.3,5.95,4.0)
text(sl,f"ECE (10 bins): {CAL['ECE_10_equal_width_bins']:.4f}  |  Brier (multiclass sum): {CAL['multiclass_brier_sum']:.4f}\nNLL: {CAL['negative_log_likelihood']:.4f}  |  Error AUROC: {CAL['error_detection_AUROC']:.4f}",.8,5.55,11.8,1.0,21)

sl=prepare(15,sources=['soft_viterbi_pseudocode.txt','transition_config.json'])
bullets(sl,['The model scores each 30-second epoch independently.',
    'The decoder balances those stage probabilities with transitions estimated from training labels.',
    'A log-space dynamic program selects one sequence for each recording; it never joins subjects.'],h=2.35,size=23)
flow(sl,['Dirichlet mean\nprobabilities','Training-label\ntransition matrix','Weighted MAP\nViterbi decoder','Decoded\nhypnogram'],y=4.3,height=1.25)
banner(sl,'Smoothing may improve aggregate scores while erasing brief stages. All transitions remain possible.')

sl=prepare(16,sources=['transition_matrix.csv','transition_config.json'])
picture(sl,'transition_matrix.png',.7,1.3,8.0,5.45)
text(sl,'Training labels only\n\nPseudocount: 1 per cell\nRows sum to one\nInitialization: stage prevalence\nSeparate recording boundaries\nEmission log-weight: 0.9',8.95,1.75,3.55,4.65,20)

sl=prepare(17,f"Raw vs Smoothed Sequence — {R['test_subjects'][0]}",sources=['test_epoch_predictions.csv'])
picture(sl,'hypnogram_subject_1.png',.55,1.3,12.2,5.55)

sl=prepare(18,sources=['raw_vs_smoothed_metrics.csv','smoothing_delta.json'])
keys=[('Accuracy','accuracy'),('Balanced accuracy','balanced_accuracy'),('Macro F1','macro_f1'),('Weighted F1','weighted_f1'),("Cohen’s kappa",'cohen_kappa')]
rows=[]
for label,k in keys:
    percent=k!='cohen_kappa';rows.append([label,pct(RAW[k]) if percent else f'{RAW[k]:.4f}',pct(SM[k]) if percent else f'{SM[k]:.4f}',
        f"{100*(SM[k]-RAW[k]):+.2f} pp" if percent else f"{SM[k]-RAW[k]:+.4f}"])
table(sl,['Metric','Raw model','Soft-Viterbi','Absolute change'],rows,y=1.65,h=3.9,widths=[3.55,2.85,2.85,2.8],size=21)
improved=[label for label,k in keys if SM[k]>RAW[k]];declined=[label for label,k in keys if SM[k]<RAW[k]]
text(sl,('Improved: '+', '.join(improved)+'.')+('\nTrade-off: '+', '.join(declined)+' decreased.' if declined else '\nAll listed aggregate metrics improved; per-subject outcomes may differ.'),.85,5.85,11.7,.95,19)

sl=prepare(19,sources=['biomarker_equations.md','subject_biomarkers.csv'])
for j,(title,body) in enumerate([('Evaluated-recording SE (%)','100 × non-Wake epochs / all evaluated epochs; a recording-duration proxy.'),
    ('WASO (minutes)','Wake epochs between first and last sleep epoch × 0.5; excludes trailing Wake.'),
    ('N3 proportion (%)','100 × N3 epochs / all evaluated epochs, including Wake in the denominator.'),
    ('SFI (transitions / sleep hour)','N2/N3/REM → Wake/N1 transitions divided by hours scored as sleep; coarse stage-transition proxy.')]):
    card(sl,title,body,.7,1.4+j*1.26,11.95,1.1)
banner(sl,'No verified time-in-bed denominator; these indicators are research measures, not diagnoses.')

sl=prepare(20,sources=['risk_rules.json','risk_rules_table.csv'])
text(sl,'Flags: SE <75% · WASO >60 min · N3 <10% · SFI >10 transitions/sleep hour',.8,1.45,11.8,.85,22,NAVY,True)
table(sl,['Priority','Research category','Exact assignment rule'],[
    ['1','High Risk','≥3 flags OR SE <65% OR SFI >15/h'],['2','Mild Risk','Otherwise, at least one flag'],['3','Healthy','Otherwise, no flags']],y=2.75,h=2.6,widths=[1.65,2.9,7.5],size=21,risk_col=1)
text(sl,'Missing/nonfinite input → Unclassified in the new evaluator.\nThresholds are source-code rules; no clinical validation is established.',.85,5.65,11.7,.9,20)

sl=prepare(21,'Test Subject Sleep Health Profiles',sources=['subject_risk_profiles.csv','biomarker_validation.csv'])
rows=[]
for sid in R['test_subjects']:
    b=BIO[(BIO.subject_id==sid)&(BIO.source=='soft_viterbi')].iloc[0]
    rows.append([sid,f'{b.Sleep_Efficiency_Pct:.2f}%',f'{b.WASO_Minutes:.1f}',f'{b.N3_Deep_Pct:.2f}%',f'{b.SFI:.2f}',b.risk_category])
table(sl,['Recording','SE proxy','WASO (min)','N3 / all','SFI (/h)','Risk'],rows,y=1.8,h=3.55,widths=[1.55,1.95,2.0,1.85,1.65,3.0],size=20,risk_col=5)
agreement=int(V[V.prediction_source=='soft_viterbi'].risk_agreement.sum())
text(sl,f"Source: Soft-Viterbi labels only. Ground-truth profiles are stored separately.\nResearch category agreement with ground truth: {agreement}/4 recordings.",.8,5.75,11.8,.85,20)
banner(sl,'SFI is a stage-transition proxy. Time in bed was not directly verified.',y=6.6)

sl=prepare(22,sources=['evaluate_authoritative_run.py','final_results.json'])
flow(sl,['Four-channel\nEDF EEG','Verified\npreprocessing','30-second\nepochs','Fast SCFormer-U\nhybrid model'],y=1.6,height=1.3)
arrow(sl,11.4,3.1,.35,.45,'down')
flow(sl,['Research profile\n+ ground-truth check','SE / WASO\nN3 / SFI','Decoded\nhypnogram','Probabilities\n+ uncertainty'],y=3.9,height=1.3)
text(sl,'Second row is read right to left: probabilities → per-recording Viterbi → biomarkers → research rules.',.8,5.75,11.8,.8,19)
# Correct second-row arrows to the actual reverse direction.
for sp in sl.shapes:
    if sp.shape_type==1 and getattr(sp,'auto_shape_type',None)==MSO_SHAPE.RIGHT_ARROW and sp.top>Inches(3.8):sp.rotation=180

sl=prepare(23,'Project Progress — Verified Deliverables',sources=['final_results.json','input_repairs.json'])
card(sl,'Completed and checked','Saved checkpoint loaded; disjoint split audited; corrupted SN001 test input reconstructed; inference rerun on all four test recordings.',.7,1.5,11.95,1.5)
card(sl,'Completed and checked','Confusion matrices, class reports, uncertainty/calibration, transition decoding, real EEG plots and all four hypnograms generated from one run.',.7,3.2,11.95,1.5)
card(sl,'Research limits remain','External validation, clinical threshold review and independent diagnostic validation have not been performed.',.7,4.9,11.95,1.5)

sl=prepare(24,'Current Research Limitations',sources=['final_results.json','calibration_metrics.json'])
bullets(sl,[f"Limited cohort: 24 SN recording IDs; only four final test recordings ({RAW['test_epochs']:,} epochs).",
    'EEG-only inputs; no multimodal PSG or cross-dataset validation.',
    'Historical notebook test reuse limits claims of a newly blinded evaluation.',
    'Evidential uncertainty and reliability require stronger independent validation.',
    'Research biomarker definitions and thresholds are not clinically validated.',
    'Training hardware and final global training RNG state were not preserved.'],size=22,h=5.25)

sl=prepare(25,'Future Work',sources=['verified scope and limitations'])
items=[('External validation','Evaluate on independent HMC recordings and another sleep dataset.'),
       ('Multimodal PSG','Test EEG + EOG + EMG with a preregistered protocol.'),
       ('Generalization','Study montage, device and population changes.'),
       ('Calibration','Fit any calibrator on validation data, then test once.'),
       ('Clinical review','Review biomarker denominators and research thresholds with experts.'),
       ('Prospective validation','Assess blinded recording-level outcomes before clinical claims.')]
for j,(title,body) in enumerate(items):card(sl,title,body,.7+(j%2)*6.15,1.45+(j//2)*1.72,5.8,1.5)

sl=prepare(26,'Verified Example System Output',sources=['eeg_example_metadata.json','subject_risk_profiles.csv'])
picture(sl,'raw_eeg_example.png',.6,1.5,6.65,3.55)
sid=R['test_subjects'][0];b=BIO[(BIO.subject_id==sid)&(BIO.source=='soft_viterbi')].iloc[0]
text(sl,f"Test recording {sid}\nSE proxy: {b.Sleep_Efficiency_Pct:.2f}%\nWASO: {b.WASO_Minutes:.1f} min\nN3 / all epochs: {b.N3_Deep_Pct:.2f}%\nSFI: {b.SFI:.2f} /sleep h\nResearch category: {b.risk_category}",7.55,1.75,5.0,3.85,22)
banner(sl,'Outputs: hypnogram, epoch probabilities, uncertainty and ground-truth validation.')

sl=prepare(27,sources=['final_results.json','overall_metrics.json'])
bullets(sl,['A compact hybrid CNN–self-attention–BiLSTM pipeline supports five-stage EEG scoring and evidential uncertainty.',
    f"On {RAW['test_epochs']:,} epochs from four test recordings, raw accuracy was {pct(RAW['accuracy'])}; Soft-Viterbi accuracy was {pct(SM['accuracy'])}.",
    'All final results use one checkpoint and the same verified test epochs; SN001’s shortened input was corrected.',
    'Research sleep profiles are reproducible, but clinical validity and external generalization remain unestablished.'],size=24,h=4.9)
banner(sl,'EEG → stage probabilities + uncertainty → decoded hypnogram → research sleep profile')

# Preserve closing design while repairing narrow original text boxes.
for sp in prs.slides[27].shapes:
    if sp.has_text_frame and sp.text.strip():
        sp.width+=Inches(.25);sp.height+=Inches(.08)
        if sp.text.strip()=='Thank You':
            sp.text_frame.word_wrap=False;sp.height=Inches(1.0)
prs.slides[27].notes_slide.notes_text_frame.text='Names and guide retained from original deck and user metadata. '+RUN
manifest.append({'slide':28,'title':'Thank You','sources':['user-supplied project metadata']})
sl=prepare(29,'Appendix',sources=['verified artifact package'])
text(sl,'Technical details and reproducible evidence',.85,2.25,11.7,.9,34,'FFFFFF',True)
text(sl,'Checkpoint architecture · classification · uncertainty · sequence decoding\nTraining settings · source equations · additional real hypnograms · literature',.85,3.5,11.7,1.4,23,'B6DCDC')
for sp in sl.shapes:
    if sp.has_text_frame and sp.text=='Appendix':
        for p in sp.text_frame.paragraphs:p.font.color.rgb=rgb('FFFFFF')

sl=prepare(30,'Appendix — Model Summary',sources=['model_layers.csv','model_summary.txt'])
rows=[['Input + permute','4 × 3000 → 3000 × 4','0'],['Conv1D 64 + BN + ReLU','750 × 64','2,112'],
    ['Conv1D 128 + BN + ReLU','188 × 128','41,600'],['Multi-head attention','188 × 128; 4 heads × key16','33,088'],
    ['Residual/LN + FFN + residual/LN','188 × 128; GELU 128→128','33,536'],['Bidirectional LSTM','188 × 128; 64 per direction','98,816'],
    ['Global average pool + dropout','128; dropout .2','0'],['Dense + dropout','128; ReLU; dropout .1','16,512'],['Evidential dense','5; softplus','645']]
table(sl,['Component','Output / setting','Parameters'],rows,y=1.45,h=4.75,widths=[4.6,5.6,1.85],size=16)
banner(sl,f"Total {HP['parameters']:,}; trainable {HP['trainable_parameters']:,}; non-trainable {HP['parameters']-HP['trainable_parameters']:,}. Full layer export in model_layers.csv.")

sl=prepare(31,'Appendix — Classification Report',sources=['classification_report.csv'])
rows=[]
for n in ['Wake','N1','N2','N3','REM','macro avg','weighted avg']:
    q=REPORT.loc[n];rows.append([n,f'{q.precision:.3f}',f'{q.recall:.3f}',f"{q['f1-score']:.3f}",f'{int(q.support):,}'])
table(sl,['Stage / average','Precision','Recall','F1','Support'],rows,y=1.5,h=4.7,widths=[3.15,2.15,2.15,2.15,2.45],size=21)
banner(sl,f"Raw model; support total {RAW['test_epochs']:,}. Macro averages weight stages equally; weighted averages use support.")

sl=prepare(32,'Appendix — Raw Confusion Counts',sources=['confusion_matrix_counts.csv'])
picture(sl,'confusion_matrix_counts.png',1.8,1.25,9.65,5.6)
sl=prepare(33,'Appendix — Soft-Viterbi Confusion Matrix',sources=['smoothed_confusion_matrix_normalized.csv'])
picture(sl,'smoothed_confusion_matrix_normalized.png',1.7,1.25,9.9,5.5)
sl.notes_slide.notes_text_frame.text+='\nRenamed unsupported self-calibrated appendix condition: no such checkpoint exists in the final project model directory.'

sl=prepare(34,'Appendix — Risk–Coverage Curve',sources=['risk_coverage_data.csv','calibration_metrics.json'])
picture(sl,'risk_coverage_curve.png',.7,1.35,11.95,5.35)
sl=prepare(35,'Appendix — Reliability and Calibration',sources=['calibration_bins.csv','calibration_metrics.json'])
picture(sl,'reliability_diagram.png',.6,1.25,8.0,5.55)
text(sl,f"Raw probabilities\n\n10 equal-width bins\nECE: {CAL['ECE_10_equal_width_bins']:.4f}\nBrier: {CAL['multiclass_brier_sum']:.4f}\nNLL: {CAL['negative_log_likelihood']:.4f}\n\nBin counts shown below the curve.",8.9,1.8,3.5,4.7,21)

sl=prepare(36,'Appendix — Training Class Weights',sources=['class_weights.csv'])
table(sl,['Stage','Training epochs','Balanced weight'],[[row.stage,f'{int(row.training_count):,}',f'{row.weight:.6f}'] for row in CW.itertuples()],y=1.7,h=3.85,widths=[3.45,4.5,4.1],size=23)
text(sl,'wk = Ntraining / (5 × Nk)\nBalanced weights use training labels only; passed to model.fit(class_weight=...).',.85,5.7,11.7,1.1,21)

sl=prepare(37,'Appendix — Final Hyperparameters',sources=['hyperparameters.json','checkpoint_config.json'])
rows=[['Optimizer / learning rate','Adam / 1 × 10⁻⁴'],['Batch / epoch limit / completed','128 / 10 / 10'],
    ['Checkpoint selection','Maximum validation evidential accuracy; epoch 9'],['Early stopping / scheduler','Patience 5, restore best; not triggered / no scheduler'],
    ['Dropout / evidence penalty','0.2 and 0.1 / coefficient 0.003'],['Label smoothing / class weighting','0.05 / balanced training-label weights'],
    ['Seed provenance','Split seed 42; final training RNG state not recoverable'],['Hardware / versions','Training hardware not recorded; checkpoint Keras 3.13.2']]
table(sl,['Setting','Verified configuration'],rows,y=1.42,h=4.95,widths=[4.25,7.8],size=18)

sl=prepare(38,'Appendix — Implemented Biomarker Equations',sources=['biomarker_equations.md'])
table(sl,['Measure','Exact source equation'],[
    ['SE proxy (%)','100 × nsleep / N'],['WASO (min)','0.5 × Wake count between first and last sleep'],['N3 / all (%)','100 × nN3 / N'],
    ['SFI (/sleep h)','Count({N2,N3,REM} → {Wake,N1}) / (nsleep / 120)']],y=1.55,h=3.15,widths=[3.4,8.65],size=21)
text(sl,'N: evaluated 30-second epochs; nsleep: non-Wake epochs; nN3: N3 epochs.\nWASO excludes trailing Wake; no-sleep WASO/SFI return 0 in source.\nNo verified time-in-bed denominator. SFI is a stage-transition proxy, not micro-arousal scoring.',.85,5.08,11.7,1.4,19)

sl=prepare(39,'Appendix — Weighted Viterbi Details',sources=['soft_viterbi_pseudocode.txt','transition_config.json'])
text(sl,'Aij = (1 + training transitions i→j) / row sum\nπj = (1 + training countj) / (5 + Ntraining)\nptj = (etj + 1) / Σk(etk + 1)',.85,1.45,11.7,1.65,23,NAVY)
text(sl,'D0j = log(πj + ε) + 0.9 log(p0j + ε)\nDtj = maxi[Dt−1,i + log(Aij + ε)] + 0.9 log(ptj + ε)\nStore argmax predecessors; backtrack from argmaxj DT−1,j.',.85,3.42,11.7,1.7,23)
text(sl,'ε = 10⁻¹²; training-only transitions; class order Wake/N1/N2/N3/REM.\nRestart at each recording or missing-epoch gap. This is weighted MAP Viterbi, not differentiable soft-DP.',.85,5.45,11.7,1.2,19)

sl=prepare(40,'Appendix — Additional Test Hypnograms',sources=['test_epoch_predictions.csv'])
picture(sl,'additional_hypnograms.png',.55,1.25,12.2,5.68)

sl=prepare(41,'Appendix — Related Literature',sources=['literature_comparison.csv','references.txt'])
rows=[['SeqSleepNet [1]\n2019','MASS; 200','EEG+EOG+EMG','20-fold subject CV','87.10% / 83.30%'],
    ['DeepSleepNet [2]\n2017','Sleep-EDF; 20','EEG Fpz–Cz','20-fold subject CV','82.00% / 76.90%'],
    ['Current raw model\n2026','HMC subset; 24','Four-channel EEG','Fixed 16/4/4 split',f"{pct(RAW['accuracy'])} / {pct(RAW['macro_f1'])}"]]
table(sl,['Study','Dataset','Modality','Protocol','Accuracy / macro F1'],rows,y=1.5,h=2.9,widths=[2.55,2.2,2.35,2.7,2.25],size=17)
text(sl,'Context only: different datasets, modalities and split protocols prevent direct ranking.\n[1] DOI: 10.1109/TNSRE.2019.2896659\n[2] DOI: 10.1109/TNSRE.2017.2721116\nHMC dataset: DOI 10.13026/t79q-fr32. Full sources and protocols in references.txt.',.85,4.85,11.7,1.75,18)

out=PROJECT/'final_sleep_stage_project_presentation.pptx';prs.save(out)
(ROOT/'slide_source_manifest.json').write_text(json.dumps(manifest,indent=2))
check=Presentation(out);assert len(check.slides)==41
pattern=re.compile(r'\[(?:INSERT|DATA REQUIRED|TO BE VERIFIED|COUNT|F1|YES/NO|VERIFIED RESULT REQUIRED|VERIFY)[^\]]*\]|PROVISIONAL',re.I)
issues=[]
for i,sl in enumerate(check.slides,1):
    for sp in sl.shapes:
        strings=[sp.text] if sp.has_text_frame else []
        if sp.has_table:strings.extend(c.text for row in sp.table.rows for c in row.cells)
        for t in strings:
            assert not pattern.search(t),(i,t)
        if sp.left<0 or sp.top<0 or sp.left+sp.width>check.slide_width+1000 or sp.top+sp.height>check.slide_height+1000:
            issues.append({'slide':i,'shape':sp.name,'issue':'outside slide'})
assert not issues,issues
print('Saved',out,'with',len(check.slides),'slides; placeholder and bounds checks passed.')
