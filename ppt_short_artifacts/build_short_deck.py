"""Edit the supplied 41-slide deck in place in memory; write a separate final copy."""
from pathlib import Path
import json, hashlib, re, math
import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

PROJECT=Path(__file__).resolve().parent.parent; ROOT=PROJECT/'ppt_artifacts'
R=json.loads((ROOT/'final_results.json').read_text()); HP=json.loads((ROOT/'hyperparameters.json').read_text())
S=json.loads((ROOT/'split_summary.json').read_text()); CAL=json.loads((ROOT/'calibration_metrics.json').read_text())
RAW=R['raw_metrics']; SM=R['smoothed_metrics']; RUN=R['run_id']
REPORT=pd.read_csv(ROOT/'classification_report.csv',index_col=0)
BIO=pd.read_csv(ROOT/'subject_biomarkers.csv'); V=pd.read_csv(ROOT/'biomarker_validation.csv')
CM=pd.read_csv(ROOT/'confusion_matrix_normalized.csv',index_col=0)
LAY=pd.read_csv(ROOT/'model_layers.csv'); CW=pd.read_csv(ROOT/'class_weights.csv')
NAVY='0B2E4F';TEAL='13A6A1';BLUE='2F6FED';GRAY='60788A';LIGHT='EEF5F8';INK='19364D'
prs=Presentation(str(PROJECT/'mid.original_backup.pptx'))
assert len(prs.slides)==41
W=prs.slide_width/Inches(1);H=prs.slide_height/Inches(1)
manifest=[]
def rgb(h):return RGBColor.from_string(h)
def box(sl,x,y,w,h,fill=LIGHT,line=None,rounded=False):
    sp=sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE,
        Inches(x),Inches(y),Inches(w),Inches(h))
    sp.fill.solid();sp.fill.fore_color.rgb=rgb(fill)
    if line:sp.line.color.rgb=rgb(line)
    else:sp.line.fill.background()
    return sp
def text(sl,txt,x,y,w,h,size=19,color=INK,bold=False,align=PP_ALIGN.LEFT,fill=None):
    sp=box(sl,x,y,w,h,fill) if fill else sl.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=sp.text_frame;tf.clear();tf.word_wrap=True
    tf.margin_left=Inches(.10 if fill else .02);tf.margin_right=Inches(.10 if fill else .02)
    tf.margin_top=Inches(.05);tf.margin_bottom=Inches(.04);tf.vertical_anchor=MSO_ANCHOR.TOP
    for i,line in enumerate(str(txt).split('\n')):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph();p.text=line;p.alignment=align
        p.space_after=Pt(9 if '\n' in str(txt) else 0)
        p.font.name='Arial';p.font.size=Pt(size);p.font.color.rgb=rgb(color);p.font.bold=bold
    return sp
def picture(sl,name,x,y,w,h):
    p=ROOT/name
    with Image.open(p) as im:iw,ih=im.size
    ratio=min(w/iw,h/ih);dw=iw*ratio;dh=ih*ratio
    return sl.shapes.add_picture(str(p),Inches(x+(w-dw)/2),Inches(y+(h-dh)/2),width=Inches(dw),height=Inches(dh))
def table(sl,headers,rows,x=.65,y=1.55,w=12.05,h=None,widths=None,size=18,risk_col=None):
    h=h or .58*(len(rows)+1)
    tbl=sl.shapes.add_table(len(rows)+1,len(headers),Inches(x),Inches(y),Inches(w),Inches(h)).table
    if widths:
        for c,width in zip(tbl.columns,widths):c.width=Inches(width)
    for i,row in enumerate([headers]+list(rows)):
        for j,value in enumerate(row):
            cell=tbl.cell(i,j);cell.text=str(value);cell.margin_left=Inches(.09);cell.margin_right=Inches(.08)
            cell.margin_top=Inches(.09);cell.margin_bottom=Inches(.06);cell.vertical_anchor=MSO_ANCHOR.MIDDLE
            cell.fill.solid();cell.fill.fore_color.rgb=rgb(NAVY if i==0 else ('FFFFFF' if i%2 else 'F1F6F8'))
            if i and j==risk_col:
                cell.fill.fore_color.rgb=rgb({'Healthy':'DBF1E9','Mild Risk':'FFF0CF','High Risk':'FADADA'}.get(str(value),'EEF5F8'))
            for p in cell.text_frame.paragraphs:
                p.font.name='Arial';p.font.size=Pt(size);p.font.bold=(i==0);p.font.color.rgb=rgb('FFFFFF' if i==0 else INK)
    return tbl
def bullets(sl,items,x=.8,y=1.45,w=11.8,h=4.9,size=22):
    return text(sl,'\n'.join('• '+str(t) for t in items),x,y,w,h,size)
def banner(sl,txt,y=6.42,color=LIGHT):text(sl,txt,.7,y,11.95,.47,16,fill=color)
def card(sl,title,body,x,y,w,h=1.7):
    box(sl,x,y,w,h,LIGHT,rounded=True);box(sl,x,y,.05,h,TEAL)
    text(sl,title,x+.17,y+.12,w-.34,.4,21,NAVY,True)
    text(sl,body,x+.17,y+.50,w-.34,h-.54,18)
def arrow(sl,x,y,w=.4,h=.3,direction='right'):
    sp=sl.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW if direction=='right' else MSO_SHAPE.DOWN_ARROW,
        Inches(x),Inches(y),Inches(w),Inches(h));sp.fill.solid();sp.fill.fore_color.rgb=rgb(BLUE);sp.line.fill.background()
def flow(sl,labels,y=2.4,height=1.35):
    n=len(labels);gap=.4;bw=(12.05-gap*(n-1))/n
    for i,label in enumerate(labels):
        x=.65+i*(bw+gap);box(sl,x,y,bw,height,LIGHT,rounded=True)
        text(sl,label,x+.06,y+.20,bw-.12,height-.3,18,NAVY,True,PP_ALIGN.CENTER)
        if i<n-1:arrow(sl,x+bw+.06,y+height/2-.12,.29,.24)
def prepare(i,title=None,sources=()):
    sl=prs.slides[i-1]
    # Keep the original full-slide background and the existing title element.
    title_sp=next((sp for sp in sl.shapes if sp.has_text_frame and sp.text.strip()),None)
    keep=[]
    for sp in list(sl.shapes):
        full=(sp.left==0 and sp.top==0 and sp.width>=prs.slide_width-1000 and sp.height==prs.slide_height)
        if full or (title_sp is not None and sp._element is title_sp._element):keep.append(sp)
        else:sp._element.getparent().remove(sp._element)
    if title_sp is None:title_sp=text(sl,title or '',.65,.4,12,.6)
    title=title or title_sp.text
    title_sp.left=Inches(.65);title_sp.top=Inches(.38);title_sp.width=Inches(12);title_sp.height=Inches(.7)
    tf=title_sp.text_frame;tf.clear();tf.word_wrap=False;tf.margin_left=0;tf.margin_top=0;tf.margin_bottom=0
    p=tf.paragraphs[0];p.text=title;p.font.name='Arial';p.font.size=Pt(28 if len(title)<66 else 25)
    p.font.color.rgb=rgb(NAVY);p.font.bold=True
    box(sl,.65,1.13,12.05,.035,TEAL)
    footer='Source: '+RUN+' | '+', '.join(sources) if sources else 'Research system — not a medical diagnosis'
    text(sl,footer,.68,7.05,11.8,.25,10.5,GRAY)
    text(sl,f'{i:02d}',12.45,7.02,.4,.28,11,GRAY,align=PP_ALIGN.RIGHT)
    sl.notes_slide.notes_text_frame.text=f"{title}\nRun: {RUN}\nSources: {', '.join(sources)}\nTest recordings: {', '.join(R['test_subjects'])}.\nCheckpoint: {R['checkpoint']}\nCheckpoint SHA256: {R['checkpoint_sha256']}\nAll classification results use identical corrected test epochs.\n"+'\n'.join(R['limitations'])
    manifest.append({'slide':i,'title':title,'sources':list(sources),'run_id':RUN})
    return sl
def pct(v):return f'{v*100:.2f}%'


# Work only on a separate copy of the completed source deck.
prs=Presentation(str(PROJECT/'ppt_short_artifacts/source_copy.pptx'))
selected=[1,2,3,41,4,7,11,12,17,21,25,27]
slide_ids=list(prs.slides._sldIdLst)
for el in list(prs.slides._sldIdLst):prs.slides._sldIdLst.remove(el)
for index in selected:prs.slides._sldIdLst.append(slide_ids[index-1])
for index,el in enumerate(slide_ids,1):
    if index not in selected:prs.part.drop_rel(el.rId)
manifest=[]
# Sources and plotted assets continue to refer to the unchanged full evaluation.
sl=prepare(2,'Introduction and Motivation',sources=['project scope','final_results.json'])
card(sl,'The engineering problem','Score EEG into Wake, N1, N2, N3 and REM, then turn the sequence into an interpretable sleep profile.',.7,1.45,11.95,1.45)
card(sl,'Why uncertainty and sequence context?','Epoch predictions can fluctuate. Evidence exposes ambiguity; training-derived transitions support coherent decoding.',.7,3.05,11.95,1.55)
card(sl,'Why personalized profiling?','Recording-level indicators summarize sleep patterns beyond a single accuracy score; their research validity must be checked against labels.',.7,4.75,11.95,1.55)
banner(sl,'Academic research system; sleep profiles are not a medical diagnosis.')
sl=prepare(3,'Project Objectives',sources=['checkpoint_config.json','transition_config.json','biomarker_validation.csv'])
for j,(a,b) in enumerate([('1  Five-stage classification','Build a compact four-channel CNN–self-attention–BiLSTM classifier.'),('2  Uncertainty and temporal refinement','Use evidential probabilities and uncertainty; decode each recording independently with training-only transitions.'),('3  Verified research profiles','Compute source-defined sleep indicators and compare predicted profiles with ground truth.')]):card(sl,a,b,.7,1.45+j*1.65,11.95,1.45)
# Existing literature, dataset and architecture slides retained, concise headings updated.
for i,title in [(4,'Literature Study'),(5,'Dataset and Preprocessing'),(6,'Current Model and Pipeline')]:
    sl=prs.slides[i-1]
    sp=next(s for s in sl.shapes if s.has_text_frame and s.top<Inches(1) and s.text.strip())
    sp.text_frame.paragraphs[0].runs[0].text=title
    for run in sp.text_frame.paragraphs[0].runs[1:]:run.text=''
# Dataset slide: retain essential audited split; add preprocessing note.
sl=prs.slides[4]
for sp in list(sl.shapes):
    if sp.has_text_frame and ('Total evaluated usable epochs:' in sp.text or 'SN001 correction:' in sp.text):sp._element.getparent().remove(sp._element)
text(sl,'22,608 usable epochs • 30 s/epoch • 256→100 Hz\n50 Hz notch • 0.3–35 Hz bandpass • per-channel epoch standardization',.8,3.03,11.7,1.12,21)
text(sl,'Recording IDs are disjoint. SN001 test input was reconstructed from real EDF.\nSN IDs are treated as subjects by the repository; person mapping is unavailable.',.8,5.45,11.7,1.0,18)
sl=prepare(7,'Results Obtained — Raw vs Soft-Viterbi',sources=['raw_vs_smoothed_metrics.csv','final_results.json'])
rows=[]
for label,key in [('Accuracy','accuracy'),('Balanced accuracy','balanced_accuracy'),('Macro F1','macro_f1'),('Weighted F1','weighted_f1'),("Cohen’s kappa",'cohen_kappa')]:
    rows.append([label,f"{RAW[key]:.4f}" if key=='cohen_kappa' else pct(RAW[key]),f"{SM[key]:.4f}" if key=='cohen_kappa' else pct(SM[key])])
table(sl,['Metric','Raw model','Soft-Viterbi'],rows,y=1.55,h=3.85,widths=[4.85,3.6,3.6],size=24)
text(sl,'Same 3,773 test epochs from four held-out recording IDs.\nAccuracy gain: +8.24 percentage points; every listed aggregate metric improved.',.85,5.7,11.7,.95,22)
sl=prepare(8,'Results — Stage Errors and Uncertainty',sources=['classification_report.csv','calibration_metrics.json'])
picture(sl,'confusion_matrix_normalized.png',.6,1.35,7.0,5.4)
text(sl,'Raw five-stage evaluation\n\nLowest F1: N1 = 0.366\nLowest recall: N2 = 44.00%\n\nError-detection AUROC: 0.580\nECE (10 bins): 0.0705\n\nUncertainty separates errors weakly; reliable referral is future work.',7.9,1.65,4.8,5.2,22)
sl=prs.slides[8]
# Retained real hypnogram: concise source-derived refinement explanation.
for sp in sl.shapes:
    if sp.has_text_frame and sp.top<Inches(1) and sp.text.strip():sp.text_frame.paragraphs[0].runs[0].text='Sequence Refinement — Real Test Recording SN009'
sl.notes_slide.notes_text_frame.text+='\nDecoder: weighted MAP Viterbi; emission log weight 0.9, training-only transitions, Laplace pseudocount 1, recording boundaries reset. All transitions possible.'
sl=prs.slides[9]
for sp in sl.shapes:
    if sp.has_text_frame and sp.top<Inches(1) and sp.text.strip():sp.text_frame.paragraphs[0].runs[0].text='Personalized Research Profiles'
sl=prepare(11,'Future Work — Multistage, Multi-model Pipeline',sources=['proposed design; not implemented','multistage_pipeline_proposal.md'])
flow(sl,['Signal quality\nchecks + flags','Model A: CNN\nWake / REM / NREM','Model B: CNN\nN1 / N2 / N3'],y=1.45,height=1.15)
text(sl,'Soft composition: P(Nk) = P(NREM) × P(Nk | NREM), k = 1,2,3.\nRetain Wake and REM probabilities; avoid irreversible hard routing.',.85,2.85,11.7,.95,20)
flow(sl,['Research profile\n+ review flag','Model C: TCN\nsequence refinement','Validation-fitted\nfusion + calibration'],y=4.03,height=1.15)
arrow(sl,10.7,3.6,.35,.35,'down')
for sp in sl.shapes:
    if sp.shape_type==1 and getattr(sp,'auto_shape_type',None)==MSO_SHAPE.RIGHT_ARROW and sp.top>Inches(4):sp.rotation=180
text(sl,'Fuse hierarchical scores with the current five-stage model as a parallel baseline.\nAdd the sequence model only after measuring whether the experts help.',.85,5.62,11.7,.95,20)
banner(sl,'Evaluate stage recall, macro F1, calibration and cost; no future performance is claimed.',y=6.65)
sl=prepare(12,'Conclusion and Future Priorities',sources=['final_results.json','multistage_pipeline_proposal.md'])
card(sl,'What was demonstrated','One verified hybrid model and training-derived decoding: 53.75% raw → 61.99% decoded accuracy on 3,773 test epochs.',.7,1.45,11.95,1.45)
card(sl,'Next engineering milestone','Train coarse-stage and NREM experts; validate soft fusion against the current model. Then test a learned temporal refiner as an ablation.',.7,3.08,11.95,1.5)
card(sl,'What is required before stronger claims','Use grouped validation and a new locked test cohort; seek external validation and expert review of biomarker definitions and risk thresholds.',.7,4.78,11.95,1.5)
banner(sl,'Current limits: four test recordings, historical test reuse, no clinical or external validation.')
prs.slides[10].notes_slide.notes_text_frame.text+='\nProposal details: multistage_pipeline_proposal.md. Calibrate the final temporal-model output separately; no uncertainty threshold is set. Use out-of-fold expert predictions to train the refiner; preserve grouped splits and use a genuinely new locked test cohort.'
# Renumber every source-retained slide, keep references/provenance in speaker notes.
for i,sl in enumerate(prs.slides,1):
    for sp in sl.shapes:
        if sp.has_text_frame and sp.left>Inches(12.3) and sp.top>Inches(6.9):sp.text_frame.paragraphs[0].runs[0].text=f'{i:02d}'
    sl.notes_slide.notes_text_frame.text+='\nCondensed 12-slide deck; future architecture is a proposal, not a reported experiment. Evidence run '+RUN
out=PROJECT/'sleep_stage_project_12_slides.pptx';prs.save(out)
assert len(Presentation(out).slides)==12
print(out)
