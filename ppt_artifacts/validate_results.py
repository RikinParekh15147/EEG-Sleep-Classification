"""Independent numerical and package-integrity acceptance checks."""
from pathlib import Path
import json, csv, math, hashlib, zipfile, re
import numpy as np
import pandas as pd
from pptx import Presentation

ROOT=Path(__file__).resolve().parent;PROJECT=ROOT.parent
R=json.loads((ROOT/'final_results.json').read_text());P=pd.read_csv(ROOT/'test_epoch_predictions.csv')
S=json.loads((ROOT/'split_summary.json').read_text());NAMES=R['class_order'];checks={}
def check(name,condition):
    checks[name]=bool(condition)
    if not condition:raise AssertionError(name)
ids=S['split_subject_ids'];allids=sum(ids.values(),[])
check('split_ids_disjoint',len(allids)==len(set(allids))==24)
check('test_subject_ids_match',set(P.subject_id)==set(ids['test']))
check('raw_smoothed_identical_epochs',len(P)==R['raw_metrics']['test_epochs']==R['smoothed_metrics']['test_epochs'])
check('epoch_identity_unique',not P.duplicated(['recording_id','epoch_index']).any())
for sid,sub in P.groupby('subject_id'):
    check('contiguous_30_second_test_'+sid,np.allclose(np.diff(sub.timestamp_seconds),30))
    check('epoch_indices_'+sid,np.array_equal(sub.epoch_index,np.arange(len(sub))))
prob=np.load(ROOT/'test_probabilities.npy');evidence=np.load(ROOT/'test_evidence.npy');alpha=np.load(ROOT/'test_dirichlet_alpha.npy')
check('probabilities_finite_normalized',np.isfinite(prob).all() and np.allclose(prob.sum(axis=1),1,atol=1e-12))
check('nonnegative_evidence',np.all(evidence>=0))
check('alpha_equals_evidence_plus_one',np.array_equal(alpha,evidence+1))
check('Dirichlet_probability_formula',np.allclose(prob,alpha/alpha.sum(axis=1,keepdims=True),atol=1e-12))
check('Dirichlet_uncertainty_formula',np.allclose(P.uncertainty,5/alpha.sum(axis=1),atol=1e-12))
check('raw_is_argmax',np.array_equal(P.raw_prediction,prob.argmax(axis=1)))
check('checkpoint_integrity',hashlib.sha256((ROOT/'authoritative_checkpoint.keras').read_bytes()).hexdigest()==R['checkpoint_sha256'])
for condition,column,prefix in [('raw','raw_prediction',''),('soft_viterbi','soft_viterbi_prediction','smoothed_')]:
    cm=np.zeros((5,5),np.int64);np.add.at(cm,(P.true_label.to_numpy(),P[column].to_numpy()),1)
    saved=pd.read_csv(ROOT/(prefix+'confusion_matrix_counts.csv'),index_col=0)
    check(condition+'_matrix_class_order',list(saved.index)==NAMES and list(saved.columns)==NAMES)
    check(condition+'_matrix_recomputed',np.array_equal(cm,saved.to_numpy()))
    check(condition+'_matrix_total',int(cm.sum())==len(P))
    report=pd.read_csv(ROOT/(prefix+'classification_report.csv'),index_col=0)
    support=cm.sum(axis=1);predicted=cm.sum(axis=0);tp=np.diag(cm)
    precision=np.divide(tp,predicted,out=np.zeros(5),where=predicted>0)
    recall=tp/support;f1=np.divide(2*precision*recall,precision+recall,out=np.zeros(5),where=precision+recall>0)
    check(condition+'_support_matches_matrix',np.array_equal(report.loc[NAMES,'support'],support))
    check(condition+'_precision_recomputed',np.allclose(report.loc[NAMES,'precision'],precision))
    check(condition+'_recall_recomputed',np.allclose(report.loc[NAMES,'recall'],recall))
    check(condition+'_f1_recomputed',np.allclose(report.loc[NAMES,'f1-score'],f1))
    accuracy=float(tp.sum()/len(P));expected=float(np.dot(support,predicted)/len(P)**2)
    computed={'accuracy':accuracy,'balanced_accuracy':float(recall.mean()),'macro_precision':float(precision.mean()),
        'macro_recall':float(recall.mean()),'macro_f1':float(f1.mean()),'weighted_f1':float(np.dot(f1,support)/support.sum()),
        'cohen_kappa':(accuracy-expected)/(1-expected)}
    metrics=R['raw_metrics'] if condition=='raw' else R['smoothed_metrics']
    for key,value in computed.items():check(condition+'_'+key+'_recomputed',math.isclose(value,metrics[key],abs_tol=1e-12))
    norm=pd.read_csv(ROOT/(prefix+'confusion_matrix_normalized.csv'),index_col=0).to_numpy()
    check(condition+'_normalized_matrix',np.allclose(norm,cm/support[:,None]))

labels=np.load(ROOT/'validation_source_labels.npz',allow_pickle=False)
transition_counts=np.ones((5,5));train_counts=np.zeros(5,dtype=int)
for sid in ids['train']:
    y=labels['train_'+sid];train_counts+=np.bincount(y,minlength=5);np.add.at(transition_counts,(y[:-1],y[1:]),1)
trans=transition_counts/transition_counts.sum(axis=1,keepdims=True)
saved_trans=pd.read_csv(ROOT/'transition_matrix.csv',index_col=0)
check('transition_class_order',list(saved_trans.index)==NAMES and list(saved_trans.columns)==NAMES)
check('transition_recomputed_from_training_only',np.allclose(trans,saved_trans.to_numpy(),atol=1e-12))
weights=pd.read_csv(ROOT/'class_weights.csv')
check('class_weights_training_only',np.allclose(weights.weight,train_counts.sum()/(5*train_counts)))
check('training_count_matches_log',train_counts.sum()==15216)
check('validation_count_matches_log',sum(len(labels['val_'+sid]) for sid in ids['val'])==3619)
for sid,sub in P.groupby('subject_id',sort=False):check('test_labels_source_'+sid,np.array_equal(labels['test_'+sid],sub.true_label))

# Independently recompute weighted Viterbi with a different implementation.
pi=(train_counts+1)/(train_counts.sum()+5)
for sid,sub in P.groupby('subject_id',sort=False):
    pp=sub[[f'prob_{name}' for name in NAMES]].to_numpy();n=len(pp)
    previous=np.log(pi+1e-12)+.9*np.log(pp[0]+1e-12);back=[]
    for t in range(1,n):
        current=np.empty(5);parents=np.empty(5,dtype=int)
        for j in range(5):
            scores=[previous[i]+math.log(trans[i,j]+1e-12) for i in range(5)]
            parents[j]=max(range(5),key=lambda i:scores[i]);current[j]=scores[parents[j]]+.9*math.log(pp[t,j]+1e-12)
        previous=current;back.append(parents)
    path=[int(previous.argmax())]
    for parents in reversed(back):path.append(int(parents[path[-1]]))
    path=path[::-1]
    check('independent_viterbi_'+sid,np.array_equal(path,sub.soft_viterbi_prediction))

bio=pd.read_csv(ROOT/'subject_biomarkers.csv')
for sid,sub in P.groupby('subject_id',sort=False):
    for source,column in [('ground_truth','true_label'),('raw','raw_prediction'),('soft_viterbi','soft_viterbi_prediction')]:
        y=sub[column].to_numpy();n=len(y);loc=np.flatnonzero(y!=0);sleep=float(len(loc));
        waso=np.sum(y[loc[0]:loc[-1]+1]==0)*.5 if len(loc) else 0
        transitions=sum(int(y[t] in [2,3,4] and y[t+1] in [0,1]) for t in range(n-1))
        expected={'Sleep_Efficiency_Pct':100*sleep/n,'WASO_Minutes':float(waso),
            'N3_Deep_Pct':100*float(np.sum(y==3))/n,'SFI':transitions/(sleep/120) if sleep else 0.}
        row=bio[(bio.subject_id==sid)&(bio.source==source)].iloc[0]
        for key,value in expected.items():check(f'biomarker_{sid}_{source}_{key}',math.isclose(row[key],value,abs_tol=1e-10))
        flags=[expected['Sleep_Efficiency_Pct']<75,expected['WASO_Minutes']>60,expected['N3_Deep_Pct']<10,expected['SFI']>10]
        category='High Risk' if sum(flags)>=3 or expected['Sleep_Efficiency_Pct']<65 or expected['SFI']>15 else 'Mild Risk' if any(flags) else 'Healthy'
        check(f'risk_{sid}_{source}',category==row.risk_category)

cal=json.loads((ROOT/'calibration_metrics.json').read_text());correct=(P.true_label==P.raw_prediction).to_numpy();conf=prob.max(axis=1)
bins=pd.read_csv(ROOT/'calibration_bins.csv');check('calibration_bin_total',bins['count'].sum()==len(P))
ece=sum(row['count']/len(P)*abs(row['accuracy']-row['mean_confidence']) for _,row in bins.iterrows() if row['count']>0)
check('ECE_recomputed',math.isclose(ece,cal['ECE_10_equal_width_bins'],abs_tol=1e-12))
brier=float(np.mean(np.sum((prob-np.eye(5)[P.true_label.to_numpy()])**2,axis=1)))
check('Brier_recomputed',math.isclose(brier,cal['multiclass_brier_sum'],abs_tol=1e-12))
nll=float(-np.log(prob[np.arange(len(P)),P.true_label.to_numpy()]).mean())
check('NLL_recomputed',math.isclose(nll,cal['negative_log_likelihood'],abs_tol=1e-12))
ranks=pd.Series(P.uncertainty).rank(method='average').to_numpy();pos=~correct;npos=pos.sum();nneg=correct.sum()
auc=(ranks[pos].sum()-npos*(npos+1)/2)/(npos*nneg)
check('error_AUROC_recomputed',math.isclose(auc,cal['error_detection_AUROC'],abs_tol=1e-12))
order=np.argsort(P.uncertainty.to_numpy(),kind='stable');risk=np.cumsum(~correct[order])/np.arange(1,len(P)+1)
check('AURC_recomputed',math.isclose(float(risk.mean()),cal['AURC_discrete_mean'],abs_tol=1e-12))
check('history_selects_epoch_9',pd.read_csv(ROOT/'training_history.csv').val_evidential_accuracy.idxmax()+1==R['selected_checkpoint_epoch']==9)
check('same_EEG_interval',json.loads((ROOT/'eeg_example_metadata.json').read_text())['reprocessed_vs_saved_max_abs_difference']==0)

for file,expected in json.loads((ROOT/'notebook_sha256_before.json').read_text()).items():check('preserved_notebook_'+Path(file).name,hashlib.sha256(Path(file).read_bytes()).hexdigest()==expected)
check('original_ppt_backup_exact',hashlib.sha256((PROJECT/'mid.original_backup.pptx').read_bytes()).hexdigest()==hashlib.sha256(Path('/mnt/c/Users/Rikin Parekh/Documents/minorProject/mid.pptx').read_bytes()).hexdigest())
deck=Presentation(PROJECT/'final_sleep_stage_project_presentation.pptx');check('slide_count_41',len(deck.slides)==41)
with zipfile.ZipFile(PROJECT/'final_sleep_stage_project_presentation.pptx') as z:check('pptx_zip_integrity',z.testzip() is None)
pattern=re.compile(r'\[(INSERT|DATA REQUIRED|TO BE VERIFIED|COUNT|F1|YES/NO|VERIFIED RESULT REQUIRED|VERIFY)[^\]]*\]|PROVISIONAL',re.I)
for i,slide in enumerate(deck.slides,1):
    strings=[]
    for shape in slide.shapes:
        if shape.has_text_frame:strings.append(shape.text)
        if shape.has_table:strings.extend(c.text for row in shape.table.rows for c in row.cells)
    check(f'slide_{i}_heading_present',any(shape.has_text_frame and shape.text.strip() and shape.top<deck.slide_height*(.5 if i==28 else .2) for shape in slide.shapes))
    check(f'slide_{i}_no_placeholders',not any(pattern.search(t) for t in strings))

result={'run_id':R['run_id'],'check_count':len(checks),'checks':checks,'passed':all(checks.values()),
        'person_identity_mapping':'Only SN recording IDs supplied; cross-person identity cannot be independently verified.'}
(ROOT/'numerical_validation.json').write_text(json.dumps(result,indent=2))
print('Independent validation passed:',len(checks),'checks')
